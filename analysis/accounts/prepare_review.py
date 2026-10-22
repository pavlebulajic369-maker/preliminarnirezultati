"""Build reviewed inputs. Extraction uses openpyxl/pypdf; never writes source files.

Image/text transcriptions below were visually reviewed on 2026-10-22.
Run evaluate_accounts.py for ordinary recalculation. Do not automatically bless new
posts by rerunning this preparation: new evidence needs a fresh manual review.
"""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'analysis/accounts'
YEARS = ('2020', '2022', '2023')

def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

def source(p, **kw):
    return {'file': p.relative_to(ROOT).as_posix(), 'sha256': digest(p), **kw}

KEYS20 = 'SNS SPS SRS SVM SPAS POKS UDS SPP METLA ZDRAVA SDA ZAVETNICI NARODNI_BLOK PSG DJB ADA JEDAN_OD_PET ZELENA_NOVA RUSKA MIR LEVIJATAN'.split()
KEYS22 = 'SNS SPS SVM SRS UJEDINJENI NADA ZAVETNICI SPP MORAMO DJB DVERI VOJVODJANI SDA SDS ADA KAD OTETE_BEBE ROMSKA RUSKI_SAVEZ'.split()
refs = {}
p = ROOT / 'data/2020_official/G202019004.xlsx'
w = load_workbook(p, read_only=True, data_only=True)
s = w.worksheets[0]
rows = list(s.values)
national = rows[3]
station_rows = [r for r in rows[4:] if isinstance(r[1], int)]
for col in [2, 3, 7, 8, 10] + list(range(12, 54, 2)):
    assert sum(r[col] or 0 for r in station_rows) == national[col], ('2020 total', col)
parties = {key: {'votes': national[12 + 2*i], 'label': rows[1][12 + 2*i].strip(),
                 'source_cell': f'{s.cell(4, 13 + 2*i).coordinate}',
                 'published_pct': national[13 + 2*i]}
           for i, key in enumerate(KEYS20)}
assert sum(v['votes'] for v in parties.values()) == national[10]
refs['2020'] = {
    'source': source(p, sheet=s.title, range='A4:BB4', extraction='national total; independently reconciled against station rows'),
    'registered_voters': national[2], 'voters_who_turned_out': national[3],
    'ballots_cast': national[7], 'invalid_ballots': national[8], 'valid_ballots': national[10],
    'parties': parties, 'official_published_share_denominator': 'voters_who_turned_out',
    'station_rows_checked': len(station_rows),
    'turnout_scope_without_kosovo_prisons_abroad': {
        'registered_voters': rows[4][2] + rows[5][2],
        'voters_who_turned_out': rows[4][3] + rows[5][3],
        'source_cells': 'Резултати!C5:D6',
        'scope': 'Central Serbia plus Vojvodina; excludes separately reported Kosovo, abroad and prison rows'},
    'note': '2020 published party percentages divide by D4 (voters who voted), not H4 (ballots in boxes). CeSID charts instead sum to about 100, so valid-vote shares are the main inferred comparison basis.'}
for v in parties.values():
    assert abs(v['published_pct'] - 100*v['votes']/national[3]) <= .00501

p = ROOT / 'data/2022_official/G202219002.xlsx'
w = load_workbook(p, read_only=True, data_only=True)
rows = list(w.active.values)
data = [r for r in rows[1:] if any(x is not None for x in r)]
assert len(data) == len({(r[2], r[4]) for r in data}) == 8267
assert all(all(isinstance(x, int) and x >= 0 for x in r[5:]) for r in data)
assert all(sum(r[10:]) == r[9] and r[8] + r[9] == r[7] and r[7] <= r[6] <= r[5] for r in data)
totals = [sum(r[i] for r in data) for i in range(5, 29)]
assert totals[:5] == [6502307, 3810559, 3806050, 112722, 3693328]
refs['2022'] = {
    'source': source(p, sheet='Sheet1', range='F2:AC8268', extraction='sum 8,267 unique municipality/station records; ignore final empty row'),
    'registered_voters': totals[0], 'voters_who_turned_out': totals[1], 'ballots_cast': totals[2],
    'invalid_ballots': totals[3], 'valid_ballots': totals[4], 'official_published_share_denominator': 'ballots_cast',
    'parties': {k: {'votes': totals[5+i], 'label': rows[0][10+i]} for i,k in enumerate(KEYS22)},
    'station_rows_checked': len(data),
    'independent_reconciliation_url': 'https://www.rik.parlament.gov.rs/vest/sr/324925/79-sednica-republicke-izborne-komisije.php',
    'note': 'National totals and all 19 list vote counts reconcile with RIK final report of 5 July 2022, including repeat voting.'}

p = ROOT / 'data/2023_official/G20246006.pdf'
old = json.loads((ROOT / 'analysis/2023/cesidbeograd/rik_reference.json').read_text(encoding='utf-8'))
pdf = PdfReader(p)
numeric_text = ''.join(pdf.pages[8].extract_text().split())
for v in old['parties'].values():
    assert str(v['votes']) in numeric_text
refs['2023'] = {k: old[k] for k in ['registered_voters', 'voters_who_turned_out', 'ballots_cast', 'invalid_ballots', 'valid_ballots', 'parties']}
refs['2023'].update(source=source(p, pages=[8,9], printed_pages=[6,7], tables=['1.1','1.2'],
    extraction='PDF text extracted; national totals, each of 18 vote counts and seat allocations checked visually on rendered pages; existing RIK transcription reused'),
    official_published_share_denominator='ballots_cast',
    note='Final parliamentary results including repeat voting on 30 December 2023 and 2 January 2024.')
for ref in refs.values():
    assert sum(v['votes'] for v in ref['parties'].values()) == ref['valid_ballots']
    assert ref['invalid_ballots'] + ref['valid_ballots'] == ref['ballots_cast']
    ref['top_six_lists'] = sorted(ref['parties'], key=lambda k: ref['parties'][k]['votes'], reverse=True)[:6]
save('official_reference.json', refs)

claims = []
def claim(pid, kind, values, **kw):
    claims.append({'post_id': str(pid), 'kind': kind, 'values': values, **kw})

for pid, sample, vals in [
    ('1274789282079703047',63.9,[62.4,10.7,2.1,2.2,4.1,2.9,1,.9,2.5,1.1,.6,1.6,.3,1.8,2.6,1,.7,.3,.2,.3,.8]),
    ('1274798166576189443',75.5,[62.5,10.7,2.1,2.2,4,2.9,1,1,2.4,1,.8,1.6,.2,1.8,2.6,1,.7,.3,.2,.3,.8]),
    ('1274811761041051649',87.4,[62.8,10.8,2,2.2,4.2,2.7,1,1,2.4,.9,.9,1.6,.2,1.8,2.6,1,.6,.3,.2,.3,.9]),
    ('1274817636006846465',91.1,[62.6,10.9,2,2.3,4.1,2.7,1,1,2.4,.9,.9,1.6,.2,1.8,2.6,1,.7,.3,.2,.3,.9])]:
    claim(pid, 'vote_share', dict(zip(KEYS20, vals)), sample_processed_pct=sample,
          basis='valid_ballots', basis_explicit=False,
          note='Valid-vote basis inferred from complete 21-list chart summing to approximately 100%. Last post separately states 4% invalid ballots. Alternative denominator errors are retained.',
          attribution='CeSID/Ipsos')
claim('1274817636006846465','invalid_ballot_share',{'invalid_pct':4},basis='ballots_cast',note='Text says invalid ballots 4%; denominator is assumed ballots cast.')
claim('1274810342326046721','turnout',{'turnout_pct':47.7},scope='without_kosovo_prisons_abroad',sample_processed_pct=87,
      note='Compare with matching territorial scope, not national turnout; earlier-year bars are historical context.')

chart22 = 'SNS UJEDINJENI SPS NADA MORAMO DVERI ZAVETNICI DJB SRS SDS SVM SPP OTETE_BEBE VOJVODJANI SDA KAD RUSKI_SAVEZ ROMSKA ADA'.split()
for pid, sample, vals in [
 ('1510718307967582218',62.5,[43.6,12.8,11.6,5.4,4.3,4,3.9,2.2,2.2,1.7,1.5,1,.8,.7,.5,.3,.2,.2,.1]),
 ('1510729200650629126',77.1,[43.1,13.2,11.4,5.4,4.5,3.9,3.9,2.3,2.2,1.7,1.6,1,.8,.7,.5,.3,.2,.2,.1])]:
    claim(pid,'vote_share',dict(zip(chart22,vals)),sample_processed_pct=sample,basis='ballots_cast',basis_explicit=False,
          note='Ballots-cast basis inferred from chart total below 100%; valid-normalized sensitivity also calculated.',attribution='CeSID/Ipsos')

oldclaims=json.loads((ROOT/'analysis/2023/cesidbeograd/claims.json').read_text(encoding='utf-8'))
for reviewed in oldclaims['claims'] + oldclaims['excluded_posts']:
    assert digest(ROOT/reviewed['image']) == reviewed['image_sha256']
for c in oldclaims['claims']:
    assert digest(ROOT/c['image']) == c['image_sha256']
    claim(c['post_id'], c['metric'], c['values'], sample_processed_pct=c['sample_processed_pct'], basis='ballots_cast' if c['metric']=='vote_share' else None,
          basis_explicit=False, attribution='CeSID/Ipsos', note='Reused previously reviewed transcription; image hash rechecked. Caption/chart scope overrides misleading Belgrade hashtag.')

claim('1274807242198585346','vote_share_bounds',{
    'SNS': {'lower':58,'upper':None,'lower_inclusive':False},
    'SPS': {'lower':9,'upper':11,'lower_inclusive':True},
    'SPAS': {'lower':3.5,'upper':4.5,'lower_inclusive':True}},
    basis='voters_who_turned_out',basis_explicit=False,sample_processed_pct=87,
    note='Literal published bounds; no fabricated SNS point estimate and no midpoint MAE. Valid-vote sensitivity is also shown.',attribution='CRTA')
claim('1274806574821842944','turnout_bounds',{'lower':48,'upper':49},scope='national',note='2016 comparison is historical context, outside selected benchmark years.')
claim('1274807584248209411','below_threshold',{'lists':['SRS','POKS','UDS','METLA','ZDRAVA','ZAVETNICI','NARODNI_BLOK','PSG','DJB','JEDAN_OD_PET','ZELENA_NOVA','MIR','LEVIJATAN'],'threshold_pct':3},
      note='Qualitative threshold claim only; explicitly excludes minority lists. Verify against official shares, not sample coverage.')
claim('1736473053465313303','registered_voters',{'2022':6502307,'2023':6500666},note='Known registration counts, not a preliminary-result forecast; excluded from forecast confidence.')
claim('1736505133297451187','winner',{'party':'SNS'},note='Text explicitly says SNS list received the most votes. Correct qualitative claim alone cannot establish numerical forecasting accuracy.')
claim('1736517035817886148','winner',{'party':'SNS'},scope_inferred=True,
      note='Victory statement interpreted from party slogan and press-conference photos; no numerical margin and no explicit election level. Keep as contextual corroboration, outside numerical score.')

# Every remaining saved post has a review disposition. Numeric non-result graphics
# are recorded as such so percentages of sample processed cannot become vote shares.
review_notes = {}
def note(ids, category, reason, extracted=None):
    for pid in ids.split():
        review_notes[pid]={'category':category,'reason':reason}
        if extracted is not None: review_notes[pid]['extracted']=extracted

note('1274772051711180800 1274777470181552128','process_report','Polling-place incident; official result totals cannot verify incident reports.')
note('1274814500814966786','link_only','Reply contains a link but no standalone numerical claim. Linked content is not saved.')
note('1274789123988107264 1274791203821834241','commentary','Question/comment about publication timing or turnout; no numerical result.')
note('1510719490430050312','different_election','Presidential election; supplied benchmarks concern national parliament.',{'sample_processed_pct':55.6,'candidate_pct':{'Obradovic':4.5,'Jovanovic':5.8,'Vucic':59.8,'Djurdjevic_Stamenkovski':4.4,'Ponos':17,'Stamenkovic':1.9,'Stojkovic':3,'Vacic':.9}})
note('1510731167590453266','different_election','Presidential election; supplied benchmarks concern national parliament.',{'sample_processed_pct':75.2,'candidate_pct':{'Obradovic':4.5,'Jovanovic':6,'Vucic':58.8,'Djurdjevic_Stamenkovski':4.4,'Ponos':17.8,'Stamenkovic':2.1,'Stojkovic':3.2,'Vacic':.9}})
note('1510735577217384457','different_election','Belgrade city assembly; national votes cast in Belgrade are a different election.',{'sample_processed_pct':46.3,'party_pct':{'SNS':39.1,'UJEDINJENI':20.1,'MORAMO':10.3,'SPS':7,'NADA':6.2,'ZAVETNICI':3.7,'DVERI':3.6,'SDS':2.7,'DJB':2.2,'SRS':1.7,'SAROVIC':.6,'RUSKI_SAVEZ':.5}})
note('1736503013278179328','different_election','Belgrade city assembly; no matching city benchmark supplied.',{'sample_processed_pct':42,'party_pct':dict(zip('SNS SPN NADA MI SPS ZAV_DVERI NS DJB_SDS CALE RUSKA LDP ZAPAD VACIC SPP'.split(),[38.2,35.1,5.9,5.3,4.9,2.6,2.2,1.6,.5,.4,.3,.3,.3,.1]))})
note('1736504577833181469','different_election','Belgrade city assembly seats; no matching city benchmark supplied.',{'sample_processed_pct':42,'seats':{'SNS':48,'SPN':43,'NADA':7,'MI':6,'SPS':6},'other_lists_seats':0})
note('1736520939683168370','different_election','Belgrade city assembly; no matching city benchmark supplied.',{'sample_processed_pct':77,'party_pct':dict(zip('SNS SPN NADA MI SPS ZAV_DVERI NS DJB_SDS CALE RUSKA LDP ZAPAD VACIC SPP'.split(),[38.8,34.5,6,5.3,4.9,2.6,2.1,1.6,.5,.4,.3,.3,.3,.1]))})
note('1274781738166779905','earlier_turnout','45.5% turnout by 19:00; cannot compare with final turnout as an error.',{'turnout_19h_pct':45.5})
note('1274776200074924032','process_report','Graphic reports irregularities, not vote shares.',{'observer_access_denied_pct':0,'polling_places_with_irregularities_pct':5})
note('1274818213742874624 1274821407038410752','event_photo','Reception photographs; no result chart or numerical election forecast.')
note('1510685669382889474','earlier_turnout','Turnout by 19:00, not final turnout.',{'turnout_19h_pct':55.1,'reported_error_pp':1})
note('1510698894216384513','earlier_turnout','Intraday turnout series cannot be evaluated against end-of-day total.',{'turnout_pct':{'10h':9.9,'13h':26.7,'16h':42.1,'19h':55.1},'reported_error_pp':{'10h':.4,'13h':.7,'16h':1,'19h':1}})
for pid,scope,pct,at in [
 ('1510705155351756803','presidential',45,'21:30'),('1510705160632381447','parliament',44,'21:30'),('1510705169360633863','belgrade_city',4,'21:30'),
 ('1510724054164119567','parliament',79,'22:45'),('1510724059461476358','presidential',77,'22:45'),('1510724065258004488','belgrade_city',25,'22:45')]:
    note(pid,'sample_progress','Percentage of sample received, not a party vote share.',{'scope':scope,'sample_processed_pct':pct,'sample_as_of_local':at})
note('1736466178099302406','earlier_turnout','Turnout by 19:00, not final turnout.',{'national_turnout_pct':55.5,'belgrade_parliamentary_turnout_pct':54.8,'national_error_pp':.9,'belgrade_error_pp':.4})
note('1736488963819340271','sample_progress','Sample coverage, not election outcome.',{'national_sample_pct':70,'belgrade_sample_pct':14,'sample_as_of_local':'21:40'})
note('1736503158015156295','sample_progress','Sample coverage, not election outcome.',{'national_sample_pct':85,'belgrade_sample_pct':40,'sample_as_of_local':'22:40'})
note('1510736370561875969','external_livestream','Livestream link plus press-conference photo; no numerical claim in saved text or attachment. Video not saved/transcribed.')
note('1736513680055902500','external_livestream','Livestream link only; video not saved/transcribed.')
note('1736492784155038064 1736499191289905374 1736505841623109652','campaign_graphic','Thank-you graphics and text contain no numerical result.')

claim_ids={c['post_id'] for c in claims}
records=[]
manifest={}
for year in YEARS:
    for path in sorted((ROOT/'data'/year).glob('*/posts.jsonl')):
        manifest[path.relative_to(ROOT).as_posix()]=digest(path)
        posts=[json.loads(line) for line in path.read_text(encoding='utf-8').splitlines() if line.strip()]
        for p in posts:
            pid=p['id']
            disposition=review_notes.get(pid)
            if pid in claim_ids:
                disposition={'category':'compared','reason':'See scored or supplementary claim records.'}
            if disposition is None:
                assert path.parent.name=='crtars', ('Unreviewed account',pid)
                disposition={'category':'process_or_publication_update','reason':'Read text: voting procedure, incident report, reply, link, or promise to publish later; no evaluable final national-result estimate in saved text.'}
            images=[]
            for m in p.get('media',[]):
                assert m['type']=='photo', ('Needs review of other media',pid)
                f=path.parent/m['local_file']
                recovered=not f.exists()
                if recovered:
                    f=OUT/'recovered_media'/year/path.parent.name/m['local_file']
                assert f.is_file(), ('Missing image',f)
                images.append({**source(f),'recovered_from_saved_public_url':recovered,'url':m.get('download_url',m.get('url')),'reviewed_visually':True})
            assert not p.get('missing_media_keys'), ('Missing media keys',pid)
            records.append({'post_id':pid,'account':path.parent.name,'year':year,'post_url':p['post_url'],
                            'created_at':p['created_at'],'text':p.get('full_text',p['text']),'images':images,**disposition})
lookup={p['post_id']:p for p in records}
for i,c in enumerate(claims):
    p=lookup[c['post_id']]
    c.update(claim_id=f"{c['post_id']}:{c['kind']}",account=p['account'],year=p['year'],
             extraction='visually reviewed image transcription' if p['images'] and c['kind'] in ('vote_share','seats','turnout') else 'reviewed saved text',
             confidence_scoring_eligible=c['kind'] in ('vote_share','vote_share_bounds'))
save('reviewed_claims.json',{'reviewed_on':'2026-10-22','scope':'Saved posts and photo attachments in election-night 20:00–00:00 windows; linked videos/pages excluded.',
    'input_posts_sha256':manifest,'records':records,'claims':claims,
    'note':'Transcriptions are reviewed data, not automatically inferred OCR. New or changed posts/images must be reviewed before evaluation can succeed.'})
print('Prepared',len(records),'post reviews,',sum(len(p['images']) for p in records),'image reviews,',len(claims),'comparison records.')
